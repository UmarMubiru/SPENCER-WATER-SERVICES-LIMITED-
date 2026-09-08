from io import BytesIO

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Inches

from .quotation_pdf import generate_quotation_pdf


def generate_quotation_docx(quotation, include_vat=None):
    """
    Generate a Word document by placing each finished PDF page onto its own Word page.
    
    Args:
        quotation: Quotation object with all necessary data
        include_vat: Boolean to include VAT in calculations (defaults to quotation.include_vat)
    
    Returns:
        BytesIO: Word document as bytes
    """
    # Use quotation's include_vat if not specified
    if include_vat is None:
        include_vat = quotation.include_vat if hasattr(quotation, 'include_vat') else True

    try:
        pdf_buffer = generate_quotation_pdf(quotation, include_vat=include_vat)
        pdf_bytes = pdf_buffer.getvalue()
        
        # Validate PDF bytes
        if not pdf_bytes or len(pdf_bytes) < 100:
            raise ValueError("Generated PDF is empty or too small")
        
        pdf_document = fitz.open(stream=pdf_bytes, filetype='pdf')
    except Exception as e:
        raise ValueError(f"PDF generation or opening failed: {str(e)}")

    docx = Document()
    section = docx.sections[0]
    section.top_margin = Cm(0)
    section.bottom_margin = Cm(0)
    section.left_margin = Cm(0)
    section.right_margin = Cm(0)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    try:
        page_count = len(pdf_document)
        for index, page in enumerate(pdf_document):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            page_bytes = pixmap.tobytes('png')
            page_image = BytesIO(page_bytes)

            paragraph = docx.add_paragraph()
            paragraph.alignment = 1
            paragraph.paragraph_format.space_before = 0
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.widow_control = True
            run = paragraph.add_run()
            run.add_picture(page_image, width=Inches(8.27))
    except Exception as e:
        pdf_document.close()
        raise ValueError(f"Page processing failed: {str(e)}")

    pdf_document.close()

    docx_bytes = BytesIO()
    docx.save(docx_bytes)
    docx_bytes.seek(0)
    return docx_bytes